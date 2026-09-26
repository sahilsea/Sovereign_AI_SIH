"""Local tools the agent loop (harness/agent_loop.py) may call.

NON-NEGOTIABLE DESIGN PRINCIPLES (same as agents/adk.py's gated tools):
1. No tool takes a principal, grade, tier or compartment argument. The
   caller's identity lives in ToolContext, set by deterministic code; the
   model has no vocabulary with which to ask for more access.
2. search_documents goes through harness/retrieve.py's two-axis gate:
   denied text is discarded before it can reach an observation. Withheld
   documents are reported by id only.
3. Every passage the agent has seen is remembered, so (a) its final
   answer's citations are verified against exactly those passages and (b)
   every file it writes inherits their classification.
4. File tools act only on the caller's own workspace (harness/workspace.py).
   Code runs only in the sandbox (agents/sandbox.py), never in-process.
5. Tools never raise into the loop: failures come back as {"ok": false,
   "error": ...} observations the model can react to.
"""

from __future__ import annotations

import base64
import csv
import io
import json
import os
import re
import time
from dataclasses import dataclass, field
from typing import Any, Callable, Optional, Sequence

from agents.sandbox import run_python_with_files
from contracts import Denial, Label, Passage, Principal, Tier
from deliver.agent_files import build_docx, build_pptx, build_xlsx
from harness import workspace
from harness.retrieve import retrieve
from trust.labels import inherit_label

OBSERVATION_CHARS = 2500
TEXT_EXTENSIONS = {"txt", "md", "csv", "json", "py", "log"}
IMAGE_EXTENSIONS = {"png", "jpg", "jpeg", "gif", "webp", "bmp", "tif", "tiff"}


@dataclass
class ToolContext:
    principal: Principal
    corpus: Sequence[Passage]
    agent: Any
    seen_passages: dict[tuple[str, int], Passage] = field(default_factory=dict)
    denials: dict[str, Denial] = field(default_factory=dict)
    artifacts: list[str] = field(default_factory=list)
    code_runs: list[dict] = field(default_factory=list)
    models_used: list[str] = field(default_factory=list)
    # Every number seen in the task text or any tool result: the only
    # numbers the agent may type inline into a spreadsheet (see
    # _ungrounded_numbers). Same idea as citation verification, for data.
    observed_numbers: set = field(default_factory=set)
    csv_outputs: list[str] = field(default_factory=list)
    inline_rows_warned: bool = False

    def observe_numbers(self, text: str) -> None:
        self.observed_numbers.update(numbers_in(text))

    @property
    def person_id(self) -> str:
        return self.principal.person_id

    def current_label(self) -> Label:
        labels = [p.label for p in self.seen_passages.values()]
        return inherit_label(labels) if labels else Label(tier=Tier.PUBLIC)

    def save(self, name: str, data: bytes) -> dict:
        info = workspace.write_bytes(self.person_id, name, data, label=self.current_label(), origin="agent")
        if info.name not in self.artifacts:
            self.artifacts.append(info.name)
        return {"file": info.name, "size_bytes": info.size_bytes, "classification": info.label.tier.value}


_NUMBER_RE = re.compile(r"(?<![\w.])-?\d[\d,]*(?:\.\d+)?")


def numbers_in(text: str) -> set[float]:
    out = set()
    for m in _NUMBER_RE.findall(str(text)):
        try:
            out.add(float(m.replace(",", "")))
        except ValueError:
            continue
    return out


def _grounded(value: float, observed: set[float]) -> bool:
    """A typed number is grounded if some observed number equals it, allowing
    for the model rounding it (1216.67 for 1216.6666...)."""
    if value in observed or value in (0.0, 1.0):
        return True
    text = repr(value)
    decimals = len(text.split(".")[1]) if "." in text and not text.endswith(".0") else 0
    tol = 0.5 * 10 ** (-decimals) + 1e-9
    return any(abs(value - o) <= tol for o in observed)


def _ungrounded_numbers(ctx: "ToolContext", payload: dict) -> list:
    bad = []
    for rows in payload.values():
        for row in rows or []:
            for cell in (row if isinstance(row, (list, tuple)) else [row]):
                if isinstance(cell, bool):
                    continue
                if isinstance(cell, (int, float)):
                    values = [float(cell)]
                elif isinstance(cell, str) and _NUMBER_RE.fullmatch(cell.strip()):
                    values = [float(cell.strip().replace(",", ""))]
                else:
                    continue
                bad.extend(v for v in values if not _grounded(v, ctx.observed_numbers))
    return bad


def _normalize_rows(rows: Any) -> list:
    """Accept a list of dicts (header from keys) as well as a list of lists."""
    rows = list(rows or [])
    if rows and all(isinstance(r, dict) for r in rows):
        header = list(rows[0].keys())
        return [header] + [[r.get(h, "") for h in header] for r in rows]
    return rows


def _ext(name: str) -> str:
    return name.rsplit(".", 1)[-1].lower() if "." in name else ""


def _require_ext(name: str, allowed: set[str], tool: str) -> str:
    name = workspace.safe_name(name)
    if _ext(name) not in allowed:
        raise workspace.WorkspaceError(f"{tool} needs a filename ending in {', '.join('.' + e for e in sorted(allowed))}.")
    return name


# ---------------------------------------------------------------------------
# Tools
# ---------------------------------------------------------------------------

def search_documents(ctx: ToolContext, query: str, top_k: int = 4) -> dict:
    top_k = max(1, min(int(top_k or 4), 6))
    passages, denials = retrieve(
        query=str(query), corpus=ctx.corpus, principal=ctx.principal, top_k=top_k,
        use_semantic=os.getenv("SEVERANCE_RETRIEVAL_MODE", "lexical") == "semantic",
    )
    for p in passages:
        ctx.seen_passages[(p.doc_id, p.page)] = p
    for d in denials:
        ctx.denials.setdefault(d.doc_id, d)
    return {
        "passages": [
            {"doc_id": p.doc_id, "page": p.page, "title": p.title, "text": p.text[:1400]}
            for p in passages
        ],
        "withheld_documents": [d.doc_id for d in denials],
        "note": "Quote passages verbatim in final citations. Withheld documents cannot be read.",
    }


def list_files(ctx: ToolContext) -> dict:
    return {"files": [
        {"name": f.name, "size_bytes": f.size_bytes, "origin": f.origin}
        for f in workspace.list_files(ctx.person_id)
    ][:40]}


def read_file(ctx: ToolContext, filename: str, max_chars: int = 6000) -> dict:
    name = workspace.safe_name(filename)
    data = workspace.read_bytes(ctx.person_id, name)
    ext = _ext(name)
    max_chars = max(500, min(int(max_chars or 6000), 12000))
    if ext in TEXT_EXTENSIONS:
        text = data.decode("utf-8", errors="replace")
    elif ext == "docx":
        from ingest.ephemeral import docx_text
        text = docx_text(data)
    elif ext == "pdf":
        text = _pdf_text(data)
    elif ext in ("xlsx", "xlsm"):
        return {"error": "This is a spreadsheet: use read_spreadsheet instead."}
    elif ext in IMAGE_EXTENSIONS:
        from ingest.ocr import ocr_image
        text = ocr_image(data) or "(OCR found no confident text; use analyze_image to have the vision model read it.)"
    else:
        return {"error": f"Cannot read .{ext} files as text."}
    return {"file": name, "chars": len(text), "text": text[:max_chars], "truncated": len(text) > max_chars}


def _pdf_text(data: bytes) -> str:
    """Per-page text; scanned pages are OCR'd on-device."""
    import pypdf
    from ingest.ephemeral import _pdf_page_to_png_base64
    from ingest.ocr import ocr_image
    reader = pypdf.PdfReader(io.BytesIO(data))
    parts = []
    for i, page in enumerate(reader.pages, start=1):
        text = (page.extract_text() or "").strip()
        source = "text"
        if not text:
            try:
                text = ocr_image(base64.b64decode(_pdf_page_to_png_base64(data, i - 1)))
                source = "OCR"
            except Exception:
                text = ""
        parts.append(f"--- Page {i} ({source}) ---\n{text or '(no readable text; try analyze_image)'}")
    return "\n".join(parts)


def write_file(ctx: ToolContext, filename: str, content: str) -> dict:
    name = _require_ext(filename, TEXT_EXTENSIONS, "write_file")
    return {"ok": True, **ctx.save(name, str(content).encode("utf-8"))}


def _rows_from_workspace(ctx: ToolContext, name: str, sheet: Optional[str] = None) -> tuple[list[str], list[list[str]]]:
    data = workspace.read_bytes(ctx.person_id, name)
    ext = _ext(name)
    if ext == "csv":
        rows = list(csv.reader(io.StringIO(data.decode("utf-8-sig", errors="replace"))))
        return ["CSV"], [r for r in rows if any(c.strip() for c in r)]
    if ext not in ("xlsx", "xlsm"):
        raise workspace.WorkspaceError("read_spreadsheet reads .xlsx, .xlsm or .csv files.")
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    names = wb.sheetnames
    ws = wb[sheet] if sheet and sheet in names else wb.worksheets[0]
    rows = [["" if v is None else v for v in r] for r in ws.iter_rows(values_only=True)]
    wb.close()
    rows = [r for r in rows if any(str(c).strip() for c in r)]
    # Skip the classification banner our own build_xlsx writes on row 1.
    if rows and str(rows[0][0]).startswith("CLASSIFICATION:"):
        rows = rows[1:]
    return names, rows


def read_spreadsheet(ctx: ToolContext, filename: str, sheet: Optional[str] = None, max_rows: int = 60) -> dict:
    name = workspace.safe_name(filename)
    sheets, rows = _rows_from_workspace(ctx, name, sheet)
    max_rows = max(5, min(int(max_rows or 60), 200))
    return {
        "file": name,
        "sheets": sheets,
        "total_rows": len(rows),
        "columns": rows[0] if rows else [],
        "rows": rows[1:max_rows + 1],
        "truncated": len(rows) - 1 > max_rows,
        "tip": "For totals/averages over many rows use run_python with input_files=[this file]; it becomes a .csv there.",
    }


def write_spreadsheet(ctx: ToolContext, filename: str, rows: Any = None, sheets: Any = None,
                      sheet_name: str = "Sheet1", from_csv: Optional[str] = None) -> dict:
    name = _require_ext(filename, {"xlsx"}, "write_spreadsheet")
    if from_csv:
        _, csv_rows = _rows_from_workspace(ctx, workspace.safe_name(from_csv))
        payload = {sheet_name or "Sheet1": csv_rows}
        return {"ok": True, **ctx.save(name, build_xlsx(payload, ctx.current_label()))}
    if isinstance(rows, list):
        rows = _normalize_rows(rows)
    if isinstance(sheets, dict):
        sheets = {k: _normalize_rows(v) for k, v in sheets.items()}
    if isinstance(sheets, dict) and sheets:
        payload = dict(sheets)
        # A model sometimes sends the header as `rows` and the data as
        # `sheets`: keep the header rather than silently dropping it.
        if isinstance(rows, list) and len(rows) == 1 and isinstance(rows[0], list):
            payload = {k: (v if v and list(v[0]) == rows[0] else [rows[0], *v]) for k, v in payload.items()}
    elif isinstance(rows, list):
        payload = {sheet_name or "Sheet1": rows}
    else:
        raise workspace.WorkspaceError("Give rows (a list of rows, first row = headers), sheets ({name: rows}), or from_csv.")
    if ctx.csv_outputs and not ctx.inline_rows_warned:
        # The sandbox already produced this data as a file: copying it through
        # the model is where rows get dropped or altered.
        ctx.inline_rows_warned = True
        return {"ok": False, "error": (
            f"run_python already wrote {', '.join(ctx.csv_outputs)}. Use write_spreadsheet with "
            f"from_csv='{ctx.csv_outputs[-1]}' so the data is copied exactly instead of retyped.")}
    bad = _ungrounded_numbers(ctx, payload)
    if bad:
        shown = ", ".join(f"{v:g}" for v in bad[:6])
        return {"ok": False, "error": (
            f"Rejected: the value(s) {shown} don't appear in the task or in any tool result, so they would be "
            "unverified. Compute them with run_python (read_rows the input, write_rows('result.csv', rows)) and "
            "then call write_spreadsheet with from_csv='result.csv'.")}
    return {"ok": True, **ctx.save(name, build_xlsx(payload, ctx.current_label()))}


def create_word_document(ctx: ToolContext, filename: str, title: str, content: str) -> dict:
    name = _require_ext(filename, {"docx"}, "create_word_document")
    body = str(content)
    if ctx.seen_passages:
        # Traceability: a reviewer can check every statement in the document
        # against the exact pages the agent read (the model's own prose in
        # the body is not itself citation-verified).
        body += "\n\n## Sources consulted by the agent\n" + "\n".join(
            f"- {p.title or p.doc_id} (`{p.doc_id}`, page {p.page})" for p in ctx.seen_passages.values()
        )
    return {"ok": True, **ctx.save(name, build_docx(str(title), body, ctx.current_label()))}


def create_presentation(ctx: ToolContext, filename: str, title: str, slides: Any) -> dict:
    name = _require_ext(filename, {"pptx"}, "create_presentation")
    if isinstance(slides, str):
        slides = [{"title": s.strip(), "bullets": []} for s in slides.split("\n") if s.strip()]
    if not isinstance(slides, list):
        raise workspace.WorkspaceError("slides must be a list of {\"title\": ..., \"bullets\": [...]} objects.")
    slides = [s if isinstance(s, dict) else {"title": str(s), "bullets": []} for s in slides]
    return {"ok": True, "slides": len(slides), **ctx.save(name, build_pptx(str(title), slides, ctx.current_label()))}


# Prepended to every agent script (standard library only). Small models
# reliably call a named helper; they unreliably hand-roll csv parsing and
# compare numeric strings as text ("6" > "13").
SANDBOX_PRELUDE = """\
import csv as _csv
def _num(v):
    try:
        return int(v)
    except (TypeError, ValueError):
        try:
            return float(v)
        except (TypeError, ValueError):
            return v
def read_rows(name):
    \"\"\"Rows of a CSV input file as dicts; numeric cells become int/float.\"\"\"
    import os as _os
    if name.lower().endswith(('.xlsx', '.xlsm')):
        name = name.rsplit('.', 1)[0] + '.csv'  # spreadsheets arrive as a CSV twin
    if not _os.path.exists(name):
        raise FileNotFoundError(f"{name} not found; files here: {sorted(_os.listdir('.'))}")
    with open(name, newline='', encoding='utf-8') as f:
        return [{k: _num(v) for k, v in r.items()} for r in _csv.DictReader(f)]
def write_rows(name, rows):
    \"\"\"Write a list of dicts (header from the keys) or a list of lists to CSV.\"\"\"
    rows = list(rows)
    with open(name, 'w', newline='', encoding='utf-8') as f:
        if rows and isinstance(rows[0], dict):
            w = _csv.DictWriter(f, fieldnames=list(rows[0].keys()))
            w.writeheader()
            w.writerows(rows)
        else:
            _csv.writer(f).writerows(rows)
# ---- agent code starts below ----
"""
PRELUDE_LINES = SANDBOX_PRELUDE.count("\n")


def run_python(ctx: ToolContext, code: str, input_files: Any = None) -> dict:
    """Run code in the sandbox with copies of the named workspace files."""
    files: dict[str, bytes] = {}
    notes = []
    requested = input_files if isinstance(input_files, list) else ([input_files] if input_files else [])
    requested = [workspace.safe_name(str(n)) for n in requested]
    # Workspace files the code names in a string literal are copied in too,
    # so a forgotten input_files argument doesn't cost a failed step.
    for f in workspace.list_files(ctx.person_id):
        stem = f.name.rsplit(".", 1)[0]
        if f.name not in requested and (f"'{f.name}'" in code or f'"{f.name}"' in code
                                       or f"'{stem}.csv'" in code or f'"{stem}.csv"' in code):
            requested.append(f.name)
    for name in requested[:5]:
        if not workspace.exists(ctx.person_id, name) and name.endswith(".csv"):
            # "stock.csv" only exists inside the sandbox, as the twin of stock.xlsx.
            twin = next((name[:-4] + ext for ext in (".xlsx", ".xlsm") if workspace.exists(ctx.person_id, name[:-4] + ext)), None)
            if twin:
                name = twin
        data = workspace.read_bytes(ctx.person_id, name)
        files[name] = data
        if _ext(name) in ("xlsx", "xlsm"):
            # The sandbox is standard-library only (no openpyxl), so hand it
            # a CSV export of the first sheet alongside the original.
            _, rows = _rows_from_workspace(ctx, name)
            buf = io.StringIO()
            csv.writer(buf).writerows(rows)
            csv_name = name.rsplit(".", 1)[0] + ".csv"
            files[csv_name] = buf.getvalue().encode("utf-8")
            notes.append(f"{name} is also available as {csv_name} (first sheet).")
    result, outputs = run_python_with_files(SANDBOX_PRELUDE + str(code), files)
    saved = []
    for out_name, data in outputs.items():
        try:
            saved.append(ctx.save(out_name, data)["file"])
            if out_name.lower().endswith(".csv") and out_name not in ctx.csv_outputs:
                ctx.csv_outputs.append(out_name)
        except workspace.WorkspaceError as exc:
            notes.append(f"Could not save {out_name}: {exc}")
    ctx.code_runs.append({"code": str(code), "result": result})
    stderr = result.stderr
    if stderr:
        stderr = stderr.replace('snippet.py", line ', 'snippet.py", prelude+line ')
        stderr += f"\n(line numbers include the {PRELUDE_LINES}-line read_rows/write_rows prelude)"
        if "NameError" in stderr or "FileNotFoundError" in stderr:
            available = sorted(files) or ["(none -- pass input_files)"]
            notes.append(f"Files in the script's folder: {', '.join(available)}. Load one with "
                         f"rows = read_rows('{next((n for n in available if n.endswith('.csv')), 'name.csv')}').")
    return {
        "exit_code": result.exit_code,
        "timed_out": result.timed_out,
        "stdout": result.stdout[:2000],
        "stderr": stderr[-1200:],
        "files_saved_to_workspace": saved,
        "notes": notes,
        "ok": result.exit_code == 0 and not result.timed_out,
    }


def analyze_image(ctx: ToolContext, filename: str, question: str, page: int = 1) -> dict:
    """Vision model on a workspace image, or one page of a PDF (scans, drawings, handwriting)."""
    name = workspace.safe_name(filename)
    data = workspace.read_bytes(ctx.person_id, name)
    ext = _ext(name)
    if ext == "pdf":
        from ingest.ephemeral import _pdf_page_to_png_base64
        b64 = _pdf_page_to_png_base64(data, max(0, int(page or 1) - 1))
    elif ext in IMAGE_EXTENSIONS:
        from ingest.ephemeral import _normalize_image_to_png_base64
        b64 = _normalize_image_to_png_base64(data)
    else:
        raise workspace.WorkspaceError("analyze_image works on images and PDF pages.")
    fn = getattr(ctx.agent, "analyze_image", None)
    if fn is None:
        raise workspace.WorkspaceError("The active backend has no vision model.")
    description = fn(str(question), b64)
    if getattr(ctx.agent, "last_model_used", None):
        ctx.models_used.append(ctx.agent.last_model_used)
    return {"file": name, "page": page if ext == "pdf" else None, "vision_model_says": description,
            "note": "Vision output is an unverified AI description, not a verbatim source."}


def calculate_tool(ctx: ToolContext, expression: str) -> dict:
    """Exact arithmetic with steps (harness/calculator.py); no sandbox needed."""
    from harness.calculator import CalculationError, calculate, evaluate, fmt
    try:
        calc = calculate(str(expression))
        if calc is None:
            # Already a bare expression the recogniser didn't need to rewrite.
            result, steps = evaluate(str(expression))
            calc = {"display": str(expression), "result": result, "result_text": fmt(result), "steps": steps}
    except CalculationError as exc:
        return {"error": f"Calculator: {exc}. For word problems or many rows, use run_python."}
    return {"ok": True, "expression": calc["display"], "result": calc["result"], "steps": calc["steps"]}


# name -> (function, one-line description, argument signature shown to the model)
TOOLS: dict[str, tuple[Callable[..., dict], str, str]] = {
    "calculate": (calculate_tool, "Exact calculator for arithmetic: + - * / ^ %, parentheses, sqrt, sin/cos/tan (degrees), log, ln, pi. Returns the result and each step. Use it for any single calculation.", '{"expression": str}'),
    "search_documents": (search_documents, "Search the organisation's internal knowledge base (SOPs, manuals, correspondence) you are cleared to read.", '{"query": str, "top_k": int (optional)}'),
    "list_files": (list_files, "List files in your workspace (uploads and files you created).", "{}"),
    "read_file": (read_file, "Read a text/.md/.csv/.json/.docx/.pdf file (scanned PDF pages are OCR'd on-device).", '{"filename": str}'),
    "read_spreadsheet": (read_spreadsheet, "Read rows from an .xlsx or .csv file.", '{"filename": str, "sheet": str (optional), "max_rows": int (optional)}'),
    "run_python": (run_python, "Run a Python 3 standard-library script in the isolated sandbox. input_files are copied into its folder (an .xlsx also arrives as a .csv of the same name). Built-in helpers: read_rows('x.csv') -> list of dicts with numbers converted; write_rows('out.csv', rows). Files it writes are saved to the workspace. print() every result and intermediate step. Example: rows = read_rows('stock.csv'); low = [dict(item=r['Item'], short=r['Min'] - r['Qty']) for r in rows if r['Qty'] < r['Min']]; write_rows('low.csv', low); print(low)", '{"code": str, "input_files": [str] (optional)}'),
    "analyze_image": (analyze_image, "Ask the local vision model about an image or a PDF page (scans, drawings, handwriting, photos).", '{"filename": str, "question": str, "page": int (optional, PDFs)}'),
    "write_file": (write_file, "Write a text file (.txt/.md/.csv/.json/.py).", '{"filename": str, "content": str}'),
    "write_spreadsheet": (write_spreadsheet, "Create an Excel .xlsx file.", '{"filename": str, "rows": [[header...], [row...]]} or {"filename": str, "sheets": {"Name": rows}} or {"filename": str, "from_csv": str}'),
    "create_word_document": (create_word_document, "Create a Word .docx deliverable (approval note, report, memo) from Markdown.", '{"filename": str, "title": str, "content": str (Markdown)}'),
    "create_presentation": (create_presentation, "Create a PowerPoint .pptx deck.", '{"filename": str, "title": str, "slides": [{"title": str, "bullets": [str]}]}'),
}


def call_tool(ctx: ToolContext, name: str, args: dict) -> tuple[bool, str, float]:
    """Execute one tool; returns (ok, observation JSON, seconds). Never raises."""
    start = time.monotonic()
    entry = TOOLS.get(name)
    if entry is None:
        return False, json.dumps({"error": f"Unknown tool '{name}'. Available: {', '.join(TOOLS)}."}), 0.0
    fn = entry[0]
    args = dict(args or {})
    ignored = []
    try:
        import inspect
        params = inspect.signature(fn).parameters
        # Common small-model aliases, then drop anything the tool doesn't take.
        for alias, real in (("sheet", "sheet_name"), ("file", "filename"), ("path", "filename"),
                            ("name", "filename"), ("text", "content"), ("body", "content")):
            if alias in args and alias not in params and real in params and real not in args:
                args[real] = args.pop(alias)
        for key in list(args):
            if key not in params:
                ignored.append(key)
                args.pop(key)
    except (TypeError, ValueError):
        pass
    try:
        result = fn(ctx, **args)
        if ignored:
            result["ignored_args"] = ignored
        ok = result.get("ok", "error" not in result)
        if ok:
            ctx.observe_numbers(json.dumps(result, default=str))
    except TypeError as exc:
        result, ok = {"error": f"Bad arguments for {name}: {exc}. Expected {entry[2]}."}, False
    except workspace.WorkspaceError as exc:
        result, ok = {"error": str(exc)}, False
    except Exception as exc:  # noqa: BLE001 - tools never raise into the loop
        result, ok = {"error": f"{type(exc).__name__}: {exc}"}, False
    text = json.dumps(result, default=str, ensure_ascii=False)
    if len(text) > OBSERVATION_CHARS:
        text = text[:OBSERVATION_CHARS] + '... [truncated]'
    return bool(ok), text, round(time.monotonic() - start, 3)


def tool_manual() -> str:
    return "\n".join(f"- {name}{sig and ' ' + sig}: {desc}" for name, (_, desc, sig) in TOOLS.items())
