"""Ephemeral, session-scoped upload handling for ad-hoc user files.

NON-NEGOTIABLE DESIGN PRINCIPLE:
Uploaded files are the user's OWN content, not governed corpus documents.
They are NEVER written to corpus/manifest.json, NEVER assigned a compartment
or tier, and NEVER subject to the two-axis security gate in trust/labels.py --
they exist only in memory, scoped to the uploading person_id, for the
lifetime of this server process. Answering questions about them is closer to
a personal document assistant than to the governed RTI corpus workbench.

The one exception is explicit: when the owner attaches an upload to an
AGENT task, its raw bytes are copied into that owner's private agent
workspace (harness/workspace.py) so the agent's file/spreadsheet tools can
open it. It is still never added to the governed corpus.
"""

from __future__ import annotations

import base64
import io
import uuid
from dataclasses import dataclass, field
from typing import Optional

from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from ingest.ocr import ocr_image

IMAGE_EXTENSIONS = {"png", "jpg", "jpeg", "gif", "webp", "bmp", "tif", "tiff"}
SPREADSHEET_EXTENSIONS = {"xlsx", "xlsm", "csv"}
TEXT_EXTENSIONS = {"txt", "md", "json", "py", "log"}
SUPPORTED_DESCRIPTION = (
    "images (png/jpg/gif/webp/bmp/tiff), scanned or digital .pdf, .pptx, .docx, "
    "spreadsheets (.xlsx/.csv) and text (.txt/.md/.json/.py)"
)
# Rows per text chunk when a spreadsheet is rendered for the drafting model.
_SHEET_ROWS_PER_CHUNK = 40
_TEXT_CHARS_PER_CHUNK = 3000


def _normalize_image_to_png_base64(raw_bytes: bytes) -> str:
    """Decode ANY supported image format and re-encode it as PNG before it
    is ever sent to the vision model.

    Some source formats (observed: .webp) reliably cause the local vision
    model to hang/time out even though the raw bytes are perfectly valid --
    almost certainly an image-decoding gap in that specific model's
    multimodal pipeline, not a SEVERANCE bug. Rather than chase every
    format-specific quirk downstream, normalize once, here, to the one
    format (PNG) every vision model handles reliably. Falls back to the
    original bytes unchanged if Pillow can't decode them (e.g. an already-
    corrupt upload) -- that failure should surface downstream as a real,
    honest vision-analysis error, not be silently swallowed here.
    """
    try:
        img = Image.open(io.BytesIO(raw_bytes))
        img.load()
        if img.mode not in ("RGB", "L"):
            img = img.convert("RGB")
        buf = io.BytesIO()
        img.save(buf, format="PNG")
        return base64.b64encode(buf.getvalue()).decode("ascii")
    except Exception:
        return base64.b64encode(raw_bytes).decode("ascii")


def _pdf_page_to_png_base64(pdf_bytes: bytes, page_index: int, dpi: int = 200) -> str:
    """Rasterize a single PDF page to a PNG, base64-encoded for the vision
    model. Used only for pages with no recoverable text layer (see
    parse_upload below) -- a page with real text is handled through the
    citation-verified text channel instead, never through this path."""
    import fitz  # PyMuPDF

    doc = fitz.open(stream=pdf_bytes, filetype="pdf")
    try:
        page = doc[page_index]
        pix = page.get_pixmap(dpi=dpi)
        return base64.b64encode(pix.tobytes("png")).decode("ascii")
    finally:
        doc.close()


@dataclass
class EphemeralImage:
    """One image extracted from an upload, base64-encoded for the vision model."""
    label: str
    base64_data: str
    mime_type: str


@dataclass
class EphemeralTextChunk:
    """One page/slide-worth of text extracted from an upload.

    source: "text" (embedded text layer), "ocr" (on-device OCR of a scanned
    page or photo), "sheet" (spreadsheet rows), or "docx".
    """
    label: str
    text: str
    source: str = "text"


@dataclass
class EphemeralUpload:
    upload_id: str
    filename: str
    person_id: str
    text_chunks: list[EphemeralTextChunk] = field(default_factory=list)
    images: list[EphemeralImage] = field(default_factory=list)
    # Original bytes, kept in memory only, so the owner can hand this file to
    # an agent task (copied into their workspace at that point, not before).
    raw_bytes: bytes = b""

    @property
    def ocr_pages(self) -> int:
        return sum(1 for c in self.text_chunks if c.source == "ocr")

    @property
    def is_spreadsheet(self) -> bool:
        return self.filename.rsplit(".", 1)[-1].lower() in SPREADSHEET_EXTENSIONS


# In-memory only: never persisted to disk or the database. Scoped to this
# server process, and every read is additionally scoped to the uploading
# person_id (see get_upload) so one user can never fetch another's upload_id.
_UPLOADS: dict[str, EphemeralUpload] = {}


def _guess_mime(ext: str) -> str:
    ext = ext.lower().lstrip(".")
    return {
        "jpg": "image/jpeg",
        "jpeg": "image/jpeg",
        "png": "image/png",
        "gif": "image/gif",
        "webp": "image/webp",
        "bmp": "image/bmp",
    }.get(ext, "application/octet-stream")


def parse_upload(filename: str, content: bytes, person_id: str) -> EphemeralUpload:
    """Parse an uploaded file into ephemeral text chunks and images.

    Supports:
    - Plain images (png/jpg/gif/webp/bmp): stored as a single image for the
      vision model to analyze directly.
    - .pptx: split PER SLIDE into text frames (-> text_chunks, answered by
      the text drafting model with the same citation-verification loop as
      governed content) and picture shapes (-> images, answered by the
      vision model). This is the actual "classify text vs image, per slide"
      behavior -- python-pptx's shape tree already carries that distinction
      natively, so it's read directly rather than guessed at.
    - .pdf: split PER PAGE the same way -- a page with an embedded text
      layer becomes a text_chunk, a page with none (a scanned page) is
      rasterized to an image. This is what powers the approval-note
      pipeline for scanned inspection reports (harness/approval_note.py).
    """
    upload_id = uuid.uuid4().hex
    ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""

    upload = EphemeralUpload(upload_id=upload_id, filename=filename, person_id=person_id)

    upload.raw_bytes = content

    if ext in IMAGE_EXTENSIONS:
        png_b64 = _normalize_image_to_png_base64(content)
        upload.images.append(EphemeralImage(label=filename, base64_data=png_b64, mime_type="image/png"))
        # A photo of a document/notice/drawing often carries printed text:
        # OCR it on-device so those words can be citation-verified. A photo
        # with no confident text yields "" and is left to the vision model.
        ocr_text = ocr_image(base64.b64decode(png_b64))
        if ocr_text:
            upload.text_chunks.append(EphemeralTextChunk(label=f"{filename} (OCR)", text=ocr_text, source="ocr"))
    elif ext == "pdf":
        # Scanned/mixed inspection report support: per PAGE, try the same
        # embedded-text extraction ingest/pdf.py uses for governed corpus
        # documents. A page WITH a text layer becomes a text_chunk (goes
        # through the same citation-verified drafting loop as everything
        # else -- see harness/approval_note.py). A page with NO recoverable
        # text (i.e. actually scanned) is rasterized to an image instead,
        # so it is answered by the vision model and clearly labeled
        # unverified, exactly like a plain image upload. This is the
        # per-page classification the approval-note pipeline depends on.
        import pypdf

        reader = pypdf.PdfReader(io.BytesIO(content))
        for page_idx, page in enumerate(reader.pages, start=1):
            text = (page.extract_text() or "").strip()
            if text:
                upload.text_chunks.append(EphemeralTextChunk(label=f"Page {page_idx}", text=text))
            else:
                try:
                    png_b64 = _pdf_page_to_png_base64(content, page_idx - 1)
                except Exception:
                    # PyMuPDF unavailable or a malformed page -- skip rather
                    # than fail the whole upload for one bad page.
                    continue
                # Scanned page: on-device OCR recovers printed text into the
                # citation-verified channel; the page image ALSO goes to the
                # vision model for handwriting, stamps, drawings and photos
                # that OCR can't read.
                ocr_text = ocr_image(base64.b64decode(png_b64))
                if ocr_text:
                    upload.text_chunks.append(
                        EphemeralTextChunk(label=f"Page {page_idx} (OCR)", text=ocr_text, source="ocr")
                    )
                upload.images.append(
                    EphemeralImage(label=f"Page {page_idx} (scanned)", base64_data=png_b64, mime_type="image/png")
                )
        if not upload.text_chunks and not upload.images:
            raise ValueError(f"'{filename}' has no pages that could be read (neither text nor image).")
    elif ext in SPREADSHEET_EXTENSIONS:
        for label, rows in _spreadsheet_rows(content, ext):
            for start in range(0, len(rows), _SHEET_ROWS_PER_CHUNK):
                block = rows[start:start + _SHEET_ROWS_PER_CHUNK]
                header = rows[0] if start else None
                lines = ([" | ".join(header)] if header else []) + [" | ".join(r) for r in block]
                upload.text_chunks.append(EphemeralTextChunk(
                    label=f"{label} rows {start + 1}-{start + len(block)}",
                    text="\n".join(lines),
                    source="sheet",
                ))
        if not upload.text_chunks:
            raise ValueError(f"'{filename}' contains no cell data.")
    elif ext == "docx":
        text = docx_text(content)
        for idx, chunk in enumerate(_split_text(text), start=1):
            upload.text_chunks.append(EphemeralTextChunk(label=f"Section {idx}", text=chunk, source="docx"))
        if not upload.text_chunks:
            raise ValueError(f"'{filename}' contains no text.")
    elif ext in TEXT_EXTENSIONS:
        text = content.decode("utf-8", errors="replace")
        for idx, chunk in enumerate(_split_text(text), start=1):
            upload.text_chunks.append(EphemeralTextChunk(label=f"Part {idx}", text=chunk))
        if not upload.text_chunks:
            raise ValueError(f"'{filename}' is empty.")
    elif ext == "pptx":
        prs = Presentation(io.BytesIO(content))
        for slide_idx, slide in enumerate(prs.slides, start=1):
            slide_text_parts: list[str] = []
            image_count = 0
            for shape in slide.shapes:
                if getattr(shape, "has_text_frame", False) and shape.text_frame.text.strip():
                    slide_text_parts.append(shape.text_frame.text.strip())
                if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                    image_count += 1
                    image = shape.image
                    upload.images.append(
                        EphemeralImage(
                            label=f"Slide {slide_idx}, Image {image_count}",
                            base64_data=_normalize_image_to_png_base64(image.blob),
                            mime_type="image/png",
                        )
                    )
            if slide_text_parts:
                upload.text_chunks.append(
                    EphemeralTextChunk(label=f"Slide {slide_idx}", text="\n".join(slide_text_parts))
                )
    else:
        raise ValueError(f"Unsupported file type '.{ext}'. Supported: {SUPPORTED_DESCRIPTION}.")

    _UPLOADS[upload_id] = upload
    return upload


def _cell(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _spreadsheet_rows(content: bytes, ext: str) -> list[tuple[str, list[list[str]]]]:
    """[(sheet label, rows as strings)] for .xlsx/.xlsm (cached values) or .csv."""
    if ext == "csv":
        import csv
        text = content.decode("utf-8-sig", errors="replace")
        rows = [[c.strip() for c in r] for r in csv.reader(io.StringIO(text))]
        rows = [r for r in rows if any(r)]
        return [("CSV", rows)] if rows else []
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(content), data_only=True, read_only=True)
    sheets = []
    for ws in wb.worksheets:
        rows = [[_cell(v) for v in r] for r in ws.iter_rows(values_only=True)]
        rows = [r for r in rows if any(r)]
        if rows:
            sheets.append((f"Sheet '{ws.title}'", rows))
    wb.close()
    return sheets


def docx_text(content: bytes) -> str:
    """Paragraphs and table rows of a .docx, in document order."""
    from docx import Document
    doc = Document(io.BytesIO(content))
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def _split_text(text: str) -> list[str]:
    """Split long text on paragraph boundaries into prompt-sized chunks."""
    chunks, current = [], ""
    for para in text.split("\n"):
        if current and len(current) + len(para) > _TEXT_CHARS_PER_CHUNK:
            chunks.append(current.strip())
            current = ""
        current += para + "\n"
    if current.strip():
        chunks.append(current.strip())
    return chunks


def get_upload(upload_id: str, person_id: str) -> Optional[EphemeralUpload]:
    """Retrieve a previously parsed upload, scoped to the uploading person_id."""
    upload = _UPLOADS.get(upload_id)
    if upload is None or upload.person_id != person_id:
        return None
    return upload


def discard_upload(upload_id: str, person_id: str) -> None:
    """Explicitly drop an upload from memory (e.g. user removes the attachment)."""
    upload = _UPLOADS.get(upload_id)
    if upload is not None and upload.person_id == person_id:
        _UPLOADS.pop(upload_id, None)
