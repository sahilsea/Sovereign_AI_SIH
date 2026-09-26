"""Deliverable PowerPoint (.pptx) report generation.

Mirrors deliver/docx.py's non-negotiable principles:
1. Classification is stamped on EVERY slide (footer) plus a dedicated title
   slide banner and the output filename -- the same "banner gets skimmed,
   footer survives copying, filename survives renaming" reasoning.
2. Refuses to build a deck for an abstained response.
"""

from __future__ import annotations

import io
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor as PptxRGBColor
from contracts import AskResponse, Tier

TIER_COLORS: dict[Tier, tuple[int, int, int]] = {
    Tier.SECRET: (185, 28, 28),
    Tier.CONFIDENTIAL: (194, 65, 12),
    Tier.INTERNAL: (29, 78, 216),
    Tier.PUBLIC: (4, 120, 87),
}

MAX_CHARS_PER_SLIDE = 900


def get_classification_text(response: AskResponse) -> str:
    tier_str = response.effective_label.tier.value.upper()
    comps = response.effective_label.compartments
    if comps:
        comp_str = ", ".join(sorted(c.value.upper() for c in comps))
        return f"CLASSIFICATION: {tier_str} // COMPARTMENTS: [{comp_str}]"
    return f"CLASSIFICATION: {tier_str} // UNRESTRICTED DISTRIBUTION"


def get_report_pptx_filename(response: AskResponse) -> str:
    row_id = response.ledger_row_id or "Unrecorded"
    tier_str = response.effective_label.tier.value.upper()
    return f"SEVERANCE_Report_Row{row_id}_{tier_str}.pptx"


def _add_classification_footer(slide, text: str, color: PptxRGBColor, prs: Presentation) -> None:
    box = slide.shapes.add_textbox(Inches(0.3), prs.slide_height - Inches(0.5), prs.slide_width - Inches(0.6), Inches(0.35))
    tf = box.text_frame
    tf.text = text
    run = tf.paragraphs[0].runs[0]
    run.font.size = Pt(10)
    run.font.bold = True
    run.font.color.rgb = color


def _chunk_text(text: str, max_len: int) -> list[str]:
    words = text.split()
    chunks: list[str] = []
    current: list[str] = []
    length = 0
    for w in words:
        if length + len(w) + 1 > max_len and current:
            chunks.append(" ".join(current))
            current, length = [], 0
        current.append(w)
        length += len(w) + 1
    if current:
        chunks.append(" ".join(current))
    return chunks or [""]


def build_report_pptx(response: AskResponse, question: str) -> bytes:
    """Build a classification-stamped slide deck summarizing an answered query."""
    if response.status == "abstained":
        raise ValueError("Cannot generate deliverable report for an abstained query response.")

    tier = response.effective_label.tier
    color_rgb = TIER_COLORS.get(tier, (0, 0, 0))
    color = PptxRGBColor(*color_rgb)
    banner_text = get_classification_text(response)

    prs = Presentation()
    blank_layout = prs.slide_layouts[6]

    # --- Title slide ---
    slide = prs.slides.add_slide(blank_layout)
    title_box = slide.shapes.add_textbox(Inches(0.7), Inches(1.5), Inches(8.6), Inches(1.5))
    tf = title_box.text_frame
    tf.word_wrap = True
    tf.text = "MRPL Document Intelligence Synthesis"
    tf.paragraphs[0].runs[0].font.size = Pt(32)
    tf.paragraphs[0].runs[0].font.bold = True

    q_box = slide.shapes.add_textbox(Inches(0.7), Inches(3.0), Inches(8.6), Inches(1.2))
    q_tf = q_box.text_frame
    q_tf.word_wrap = True
    q_tf.text = f"Query: {question}"

    banner_box = slide.shapes.add_textbox(Inches(0.7), Inches(4.2), Inches(8.6), Inches(0.6))
    b_tf = banner_box.text_frame
    b_tf.text = banner_text
    b_tf.paragraphs[0].runs[0].font.size = Pt(16)
    b_tf.paragraphs[0].runs[0].font.bold = True
    b_tf.paragraphs[0].runs[0].font.color.rgb = color
    _add_classification_footer(slide, banner_text, color, prs)

    # --- Findings slide(s): chunk the answer text across slides ---
    for chunk in _chunk_text(response.answer, MAX_CHARS_PER_SLIDE):
        slide = prs.slides.add_slide(blank_layout)
        head = slide.shapes.add_textbox(Inches(0.5), Inches(0.4), Inches(9), Inches(0.7))
        head.text_frame.text = "Synthesized Findings (quotations verified)"
        head.text_frame.paragraphs[0].runs[0].font.size = Pt(24)
        head.text_frame.paragraphs[0].runs[0].font.bold = True

        body = slide.shapes.add_textbox(Inches(0.5), Inches(1.2), Inches(9), Inches(5.5))
        body.text_frame.word_wrap = True
        body.text_frame.text = chunk
        body.text_frame.paragraphs[0].font.size = Pt(16)
        _add_classification_footer(slide, banner_text, color, prs)

    # --- Citations slide ---
    if response.citations:
        slide = prs.slides.add_slide(blank_layout)
        head = slide.shapes.add_textbox(Inches(0.5), Inches(0.4), Inches(9), Inches(0.7))
        head.text_frame.text = "Verbatim Citation Verification"
        head.text_frame.paragraphs[0].runs[0].font.size = Pt(24)
        head.text_frame.paragraphs[0].runs[0].font.bold = True

        rows = len(response.citations) + 1
        table_shape = slide.shapes.add_table(rows, 3, Inches(0.5), Inches(1.2), Inches(9), Inches(0.5) * rows)
        table = table_shape.table
        table.cell(0, 0).text = "Document ID"
        table.cell(0, 1).text = "Page"
        table.cell(0, 2).text = "Verbatim Quote"
        for i, cit in enumerate(response.citations, start=1):
            table.cell(i, 0).text = cit.doc_id
            table.cell(i, 1).text = str(cit.page)
            table.cell(i, 2).text = cit.quote
        _add_classification_footer(slide, banner_text, color, prs)

    # --- Withheld materials slide ---
    if response.denials:
        slide = prs.slides.add_slide(blank_layout)
        head = slide.shapes.add_textbox(Inches(0.5), Inches(0.4), Inches(9), Inches(0.7))
        head.text_frame.text = "Severability Clause & Withheld Documents"
        head.text_frame.paragraphs[0].runs[0].font.size = Pt(24)
        head.text_frame.paragraphs[0].runs[0].font.bold = True

        body = slide.shapes.add_textbox(Inches(0.5), Inches(1.2), Inches(9), Inches(5))
        body.text_frame.word_wrap = True
        lines = [f"{d.doc_id}: {d.reason}" for d in response.denials]
        body.text_frame.text = "\n".join(lines)
        _add_classification_footer(slide, banner_text, color, prs)

    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()
