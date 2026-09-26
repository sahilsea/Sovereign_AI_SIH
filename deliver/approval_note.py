"""Approval note Word document generation.

Distinct from deliver/docx.py's generic Q&A report template: this renders
the specific "approval note" layout the PS names as its example deliverable
(title, executive summary, findings, recommendation, sign-off block) rather
than a generic synthesis report. Reuses the same 3-way classification
stamping helpers so both deliverable types look and behave consistently.
"""

from __future__ import annotations

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor
from contracts import AskResponse
from deliver.docx import TIER_COLORS, get_classification_banner_text


def get_approval_note_filename(response: AskResponse) -> str:
    row_id = response.ledger_row_id or "Unrecorded"
    tier_str = response.effective_label.tier.value.upper()
    return f"SEVERANCE_ApprovalNote_Row{row_id}_{tier_str}.docx"


def build_approval_note_docx(response: AskResponse) -> bytes:
    """Build the formal approval-note Word document from an AskResponse
    whose `approval_note` field is populated (see harness/approval_note.py).
    """
    if response.approval_note is None:
        raise ValueError("build_approval_note_docx requires response.approval_note to be set.")
    note = response.approval_note

    doc = Document()
    tier = response.effective_label.tier
    stamp_color = TIER_COLORS.get(tier, RGBColor(0, 0, 0))
    banner_text = get_classification_banner_text(response)

    for section in doc.sections:
        header_p = section.header.paragraphs[0]
        header_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        header_run = header_p.add_run(banner_text)
        header_run.bold = True
        header_run.font.size = Pt(8.5)
        header_run.font.color.rgb = stamp_color

        footer_p = section.footer.paragraphs[0]
        footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        footer_run = footer_p.add_run(banner_text)
        footer_run.bold = True
        footer_run.font.size = Pt(8.5)
        footer_run.font.color.rgb = stamp_color

    banner_p = doc.add_paragraph()
    banner_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    b_run = banner_p.add_run(f"*** {banner_text} ***")
    b_run.bold = True
    b_run.font.size = Pt(13)
    b_run.font.color.rgb = stamp_color

    doc.add_heading(note.title, level=1)

    meta_p = doc.add_paragraph()
    meta_p.add_run("Source Document: ").bold = True
    meta_p.add_run(f"{note.source_filename}\n")
    meta_p.add_run("Ledger Row ID: ").bold = True
    meta_p.add_run(f"#{response.ledger_row_id or 'Unrecorded'}\n")

    doc.add_heading("Executive Summary", level=2)
    doc.add_paragraph(note.summary)

    doc.add_heading("Findings", level=2)
    if note.verified_findings:
        p = doc.add_paragraph()
        p.add_run("Verified findings (verbatim, citation-checked against source text):").italic = True
        for c in note.verified_findings:
            bullet = doc.add_paragraph(style="List Bullet")
            bullet.add_run(f'"{c.quote}" ').italic = False
            bullet.add_run(f"[{c.doc_id}, p.{c.page}]").italic = True
    if note.visual_observations:
        p = doc.add_paragraph()
        p.add_run(
            "Unverified visual observations (AI-described from scanned pages with no text layer "
            "-- not verbatim-verified):"
        ).italic = True
        for obs in note.visual_observations:
            doc.add_paragraph(obs, style="List Bullet")
    if not note.verified_findings and not note.visual_observations:
        doc.add_paragraph("No extractable findings were identified in the source document.")

    doc.add_heading("Recommendation", level=2)
    doc.add_paragraph(note.recommendation)

    doc.add_heading("Approval", level=2)
    table = doc.add_table(rows=3, cols=2)
    table.style = "Table Grid"
    rows_labels = ["Prepared By (System-Generated)", "Reviewed By", "Approved By"]
    for i, label in enumerate(rows_labels):
        table.rows[i].cells[0].text = label
        table.rows[i].cells[1].text = "" if i > 0 else "SEVERANCE Workbench (auto-drafted, human review required)"

    return _doc_to_bytes(doc)


def _doc_to_bytes(doc: Document) -> bytes:
    import io
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
