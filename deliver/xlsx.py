"""Deliverable Excel (.xlsx) report generation.

Structured, tabular counterpart to deliver/docx.py's narrative report --
same classification-stamping discipline (a header banner row on every
sheet, plus a stamped filename), same refusal to render an abstained
response, but shaped as data (one row per citation/denial) instead of prose,
for the analysts who want to pivot/filter the audit trail rather than read it.
"""

from __future__ import annotations

import io
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from contracts import AskResponse, Tier

TIER_HEX: dict[Tier, str] = {
    Tier.SECRET: "B91C1C",
    Tier.CONFIDENTIAL: "C2410C",
    Tier.INTERNAL: "1D4ED8",
    Tier.PUBLIC: "047857",
}


def get_classification_text(response: AskResponse) -> str:
    tier_str = response.effective_label.tier.value.upper()
    comps = response.effective_label.compartments
    if comps:
        comp_str = ", ".join(sorted(c.value.upper() for c in comps))
        return f"CLASSIFICATION: {tier_str} // COMPARTMENTS: [{comp_str}]"
    return f"CLASSIFICATION: {tier_str} // UNRESTRICTED DISTRIBUTION"


def get_report_xlsx_filename(response: AskResponse) -> str:
    row_id = response.ledger_row_id or "Unrecorded"
    tier_str = response.effective_label.tier.value.upper()
    return f"SEVERANCE_Report_Row{row_id}_{tier_str}.xlsx"


def _stamp_banner(ws, banner_text: str, hex_color: str, width: int) -> None:
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=width)
    cell = ws.cell(row=1, column=1, value=f"*** {banner_text} ***")
    cell.font = Font(bold=True, color="FFFFFF", size=12)
    cell.fill = PatternFill(start_color=hex_color, end_color=hex_color, fill_type="solid")
    cell.alignment = Alignment(horizontal="center")


def build_report_xlsx(response: AskResponse, question: str) -> bytes:
    """Build a classification-stamped workbook: a Summary sheet, a Citations
    sheet (one row per verified quote), and a Denials sheet."""
    if response.status == "abstained":
        raise ValueError("Cannot generate deliverable report for an abstained query response.")

    tier = response.effective_label.tier
    hex_color = TIER_HEX.get(tier, "000000")
    banner_text = get_classification_text(response)

    wb = Workbook()

    # --- Summary sheet ---
    ws = wb.active
    ws.title = "Summary"
    _stamp_banner(ws, banner_text, hex_color, width=2)
    ws.cell(row=3, column=1, value="Query").font = Font(bold=True)
    ws.cell(row=3, column=2, value=question)
    ws.cell(row=4, column=1, value="Ledger Row ID").font = Font(bold=True)
    ws.cell(row=4, column=2, value=response.ledger_row_id or "Unrecorded")
    ws.cell(row=5, column=1, value="Effective Classification").font = Font(bold=True)
    ws.cell(row=5, column=2, value=tier.value.upper())
    ws.cell(row=7, column=1, value="Answer").font = Font(bold=True)
    ws.cell(row=8, column=1, value=response.answer)
    ws.cell(row=8, column=1).alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells(start_row=8, start_column=1, end_row=8, end_column=6)
    ws.row_dimensions[8].height = 120
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 60

    # --- Citations sheet ---
    ws2 = wb.create_sheet("Citations")
    _stamp_banner(ws2, banner_text, hex_color, width=3)
    headers = ["Document ID", "Page", "Verbatim Quote"]
    for col, h in enumerate(headers, start=1):
        c = ws2.cell(row=3, column=col, value=h)
        c.font = Font(bold=True)
    for i, cit in enumerate(response.citations, start=4):
        ws2.cell(row=i, column=1, value=cit.doc_id)
        ws2.cell(row=i, column=2, value=cit.page)
        ws2.cell(row=i, column=3, value=cit.quote)
    ws2.column_dimensions["A"].width = 24
    ws2.column_dimensions["B"].width = 10
    ws2.column_dimensions["C"].width = 80

    # --- Denials sheet ---
    ws3 = wb.create_sheet("Withheld Documents")
    _stamp_banner(ws3, banner_text, hex_color, width=2)
    ws3.cell(row=3, column=1, value="Document ID").font = Font(bold=True)
    ws3.cell(row=3, column=2, value="Reason").font = Font(bold=True)
    for i, denial in enumerate(response.denials, start=4):
        ws3.cell(row=i, column=1, value=denial.doc_id)
        ws3.cell(row=i, column=2, value=denial.reason)
    ws3.column_dimensions["A"].width = 24
    ws3.column_dimensions["B"].width = 80

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
