"""Tests for the egress guard, model registry, OCR and new upload types.

Verifies:
1. The egress guard blocks public addresses and allows loopback/private.
2. The model registry auto-selects the best INSTALLED model per task, and an
   env pin wins only when that model is installed.
3. On-device OCR recovers printed text from an image-only scan.
4. Spreadsheet, CSV, DOCX and text uploads are parsed into text chunks.
5. The deterministic intent fallback sends deliverable requests to the agent.
"""

import io
import shutil
import socket

import pytest
from docx import Document
from openpyxl import Workbook
from PIL import Image, ImageDraw, ImageFont

from agents import registry
from agents.real import RealLlmAgent
from ingest.ephemeral import parse_upload
from ingest.ocr import ocr_image
from trust import egress


# --- egress guard ----------------------------------------------------------

@pytest.fixture
def guard():
    previous = egress.snapshot()["guard"]["mode"]
    egress.install_guard("block")
    yield
    egress.install_guard(previous)


def test_guard_blocks_public_address_before_connecting(guard):
    s = socket.socket(socket.AF_INET, socket.SOCK_STREAM)
    try:
        with pytest.raises(egress.EgressBlocked):
            s.connect(("1.1.1.1", 443))
    finally:
        s.close()
    blocked = egress.snapshot()["guard"]["blocked_attempts"]
    assert blocked[-1]["remote_ip"] == "1.1.1.1" and blocked[-1]["action"] == "blocked"


def test_guard_allows_loopback(guard):
    listener = socket.socket(socket.AF_INET, socket.SOCK_STREAM)
    listener.bind(("127.0.0.1", 0))
    listener.listen(1)
    client = socket.socket(socket.AF_INET, socket.SOCK_STREAM)
    try:
        client.connect(listener.getsockname())  # must not raise
    finally:
        client.close()
        listener.close()


@pytest.mark.parametrize("ip,local", [
    ("127.0.0.1", True), ("10.2.3.4", True), ("192.168.1.9", True), ("172.20.0.5", True),
    ("::1", True), ("8.8.8.8", False), ("1.1.1.1", False), ("2606:4700::1111", False),
])
def test_local_address_classification(ip, local):
    assert egress.is_local_address(ip) is local


# --- model registry --------------------------------------------------------

def test_registry_picks_best_installed_model(monkeypatch):
    for var in ("OLLAMA_MODEL", "OLLAMA_CODE_MODEL", "OLLAMA_VISION_MODEL", "OLLAMA_PLANNER_MODEL"):
        monkeypatch.delenv(var, raising=False)
    monkeypatch.setattr(registry, "installed_models",
                        lambda force=False: {"llama3.2:latest", "granite4.1:3b", "qwen2.5-coder:3b", "granite3.2-vision:latest"})
    assert registry.pick("document") == "granite4.1:3b"   # outranks llama3.2
    assert registry.pick("code") == "qwen2.5-coder:3b"
    assert registry.pick("vision") == "granite3.2-vision:latest"
    assert registry.pick("document") != registry.pick("code")  # different task -> different model


def test_registry_env_pin_wins_only_when_installed(monkeypatch):
    monkeypatch.setattr(registry, "installed_models", lambda force=False: {"llama3.2:latest", "granite4.1:3b"})
    monkeypatch.setenv("OLLAMA_MODEL", "llama3.2")
    assert registry.pick("document") == "llama3.2"
    monkeypatch.setenv("OLLAMA_MODEL", "not-pulled:7b")
    assert registry.pick("document") == "granite4.1:3b"


def test_new_model_is_selected_by_adding_a_registry_entry(monkeypatch, tmp_path):
    reg = tmp_path / "models.json"
    reg.write_text('{"models": [{"id": "future-coder:9b", "tasks": ["code"], "priority": 999},'
                   ' {"id": "qwen2.5-coder:3b", "tasks": ["code"], "priority": 70}]}')
    monkeypatch.setattr(registry, "REGISTRY_PATH", reg)
    monkeypatch.delenv("OLLAMA_CODE_MODEL", raising=False)
    monkeypatch.setattr(registry, "installed_models", lambda force=False: {"future-coder:9b", "qwen2.5-coder:3b"})
    assert registry.pick("code") == "future-coder:9b"


# --- OCR -------------------------------------------------------------------

@pytest.mark.skipif(shutil.which("tesseract") is None, reason="tesseract not installed")
def test_ocr_recovers_text_from_image_only_scan():
    img = Image.new("L", (1400, 300), 255)
    d = ImageDraw.Draw(img)
    try:
        font = ImageFont.truetype("/System/Library/Fonts/Supplemental/Arial.ttf", 40)
    except OSError:
        font = ImageFont.load_default(size=40)
    d.text((40, 40), "Pressure safety valve PSV-2104A certificate expired.", font=font, fill=0)
    d.text((40, 140), "Coating damaged over three square metres on the south side.", font=font, fill=0)
    buf = io.BytesIO()
    img.save(buf, "PNG")
    text = ocr_image(buf.getvalue())
    assert "PSV-2104A" in text
    assert "south side" in text

    upload = parse_upload("scan.png", buf.getvalue(), "p1")
    assert upload.ocr_pages == 1 and upload.images  # OCR text AND the image for the vision model


def test_ocr_returns_nothing_for_blank_image():
    buf = io.BytesIO()
    Image.new("L", (400, 200), 255).save(buf, "PNG")
    assert ocr_image(buf.getvalue()) == ""


# --- new upload types ------------------------------------------------------

def test_xlsx_upload_is_parsed_per_sheet():
    wb = Workbook()
    ws = wb.active
    ws.title = "Stock"
    ws.append(["Item", "Qty"])
    for i in range(45):
        ws.append([f"part-{i}", i])
    buf = io.BytesIO()
    wb.save(buf)
    up = parse_upload("stock.xlsx", buf.getvalue(), "p1")
    assert up.is_spreadsheet
    assert up.text_chunks[0].label == "Sheet 'Stock' rows 1-40"
    assert up.text_chunks[1].text.startswith("Item | Qty")  # header repeated on later chunks
    assert up.raw_bytes == buf.getvalue()


def test_csv_docx_and_text_uploads():
    up = parse_upload("a.csv", b"unit,flow\nCDU,120\n", "p1")
    assert "CDU | 120" in up.text_chunks[0].text

    doc = Document()
    doc.add_paragraph("Approval is sought for replacing pump P-101 impeller.")
    buf = io.BytesIO()
    doc.save(buf)
    up = parse_upload("memo.docx", buf.getvalue(), "p1")
    assert "P-101 impeller" in up.text_chunks[0].text

    up = parse_upload("notes.txt", b"Line one\nLine two", "p1")
    assert up.text_chunks[0].text == "Line one\nLine two"

    with pytest.raises(ValueError):
        parse_upload("x.exe", b"MZ", "p1")


# --- intent fallback -------------------------------------------------------

@pytest.mark.parametrize("q,expected", [
    ("make an excel file with monthly steam consumption", "task"),
    ("draft an approval note for the pump overhaul", "task"),
    ("write a python script to sort numbers", "code"),
    ("what is the gas leak procedure", "content"),
])
def test_keyword_intent_fallback(q, expected):
    assert RealLlmAgent._keyword_category(q) == expected
