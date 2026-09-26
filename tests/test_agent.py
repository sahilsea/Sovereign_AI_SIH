"""Tests for the multi-step agent loop, its tools, and the agent workspace.

Verifies:
1. Plan -> tool calls -> verified final answer, with a real deliverable file.
2. Files and answers inherit the classification of every passage seen.
3. A final answer whose citations don't verify is withheld.
4. search_documents discards denied text at the gate.
5. Workspace filenames can't escape the owner's directory.
6. run_python gets workspace files (xlsx also as csv) and saves its outputs.
7. A repeated identical tool call is refused.
"""

import io
import json
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from contracts import Compartment, Label, Passage, Principal, Tier
from harness import workspace
from harness.agent_loop import run_agent_task
from harness.agent_tools import ToolContext, call_tool
from trust.conversations import init_conversations_table
from trust.ledger import init_ledger_table, read as ledger_read
from trust.reports import init_reports_table

HSE_TEXT = "During a gas leak, isolate the source, stop all hot work and evacuate personnel upwind to the assembly point."
VIG_TEXT = "The procurement audit found irregular single-bid awards totalling forty crore rupees."


class ScriptedAgent:
    """Replays canned JSON replies in place of the planner model."""

    def __init__(self, replies):
        self.replies = [json.dumps(r) for r in replies]
        self.last_model_used = "scripted-planner"
        self.calls = []

    def chat_json(self, messages, purpose="planning"):
        self.calls.append(messages[-1]["content"])
        return self.replies.pop(0) if self.replies else json.dumps({"final_answer": "done", "citations": []})


@pytest.fixture
def env(tmp_path: Path, monkeypatch):
    monkeypatch.setenv("SEVERANCE_WORKSPACE_DIR", str(tmp_path / "ws"))
    monkeypatch.setenv("SEVERANCE_RETRIEVAL_MODE", "lexical")
    db = str(tmp_path / "agent.db")
    for init in (init_ledger_table, init_reports_table, init_conversations_table):
        init(db)
    corpus = [
        Passage(doc_id="hse-sop", page=3, title="HSE SOP", text=HSE_TEXT,
                label=Label(tier=Tier.CONFIDENTIAL, compartments=frozenset([Compartment.HSE]))),
        Passage(doc_id="vig-audit", page=1, title="Vigilance Audit", text=VIG_TEXT,
                label=Label(tier=Tier.CONFIDENTIAL, compartments=frozenset([Compartment.VIGILANCE]))),
    ]
    principal = Principal(person_id="eng-001", name="Engineer", job_title="Engineer", grade="E",
                          compartments=frozenset([Compartment.HSE]))
    return db, corpus, principal


def test_agent_plans_acts_and_delivers_verified_word_file(env):
    db, corpus, principal = env
    agent = ScriptedAgent([
        {"plan": ["1. search_documents for gas leak", "2. create_word_document"]},
        {"thought": "find the SOP", "tool": "search_documents", "args": {"query": "gas leak evacuate"}},
        {"thought": "write it", "tool": "create_word_document",
         "args": {"filename": "gas_leak_note.docx", "title": "Gas Leak Note", "content": "## Steps\n- Isolate the source"}},
        {"final_answer": "Created the note.", "citations": [
            {"doc_id": "hse-sop", "page": 3, "quote": "isolate the source, stop all hot work and evacuate personnel upwind"}]},
    ])
    res = run_agent_task("Draft a gas leak note as Word", principal, agent, db, corpus)

    assert res.status == "answered"
    assert res.plan == ["1. search_documents for gas leak", "2. create_word_document"]
    assert [s.tool for s in res.agent_trace] == ["search_documents", "create_word_document"]
    assert all(s.ok for s in res.agent_trace)
    assert [a.name for a in res.artifacts] == ["gas_leak_note.docx"]
    # Inherited from the HSE passage the agent read.
    assert res.effective_label.tier == Tier.CONFIDENTIAL
    assert Compartment.HSE in res.effective_label.compartments
    assert res.artifacts[0].label.tier == Tier.CONFIDENTIAL
    assert workspace.read_bytes("eng-001", "gas_leak_note.docx")[:2] == b"PK"  # a real .docx (zip)
    assert ledger_read(db, limit=1)[0].action == "AGENT_TASK_ANSWERED"


def test_unverifiable_final_answer_is_withheld(env):
    db, corpus, principal = env
    bad = {"final_answer": "Evacuate downwind.", "citations": [
        {"doc_id": "hse-sop", "page": 3, "quote": "evacuate all personnel downwind immediately to the gate"}]}
    agent = ScriptedAgent([
        {"plan": ["1. search"]},
        {"tool": "search_documents", "args": {"query": "gas leak"}},
        bad, bad,
    ])
    res = run_agent_task("What do we do in a gas leak?", principal, agent, db, corpus)
    assert res.status == "abstained"
    assert "withheld" in res.answer.lower()
    assert "CITATION CHECK FAILED" in agent.calls[-1]  # failure fed back before giving up


def test_search_tool_discards_denied_text(env):
    _, corpus, principal = env
    ctx = ToolContext(principal=principal, corpus=corpus, agent=None)
    ok, observation, _ = call_tool(ctx, "search_documents", {"query": "procurement audit single-bid awards"})
    assert ok
    assert "forty crore" not in observation
    assert "vig-audit" in json.loads(observation)["withheld_documents"]
    assert ("vig-audit", 1) not in ctx.seen_passages


@pytest.mark.parametrize("name", ["..", "a..b", ".meta.json", ".hidden", "", "a" * 200])
def test_workspace_rejects_unsafe_names(env, name):
    with pytest.raises(workspace.WorkspaceError):
        workspace.safe_name(name)


def test_workspace_flattens_paths_into_owner_dir(env):
    info = workspace.write_bytes("eng-001", "/etc/passwd", b"x")
    assert info.name == "passwd"
    assert (workspace.root() / "eng-001" / "passwd").is_file()
    assert not workspace.exists("someone-else", "passwd")
    assert workspace.write_bytes("eng-001", "../../escape.txt", b"y").name == "escape.txt"
    assert (workspace.root() / "eng-001" / "escape.txt").is_file()


def test_run_python_reads_spreadsheet_and_saves_output(env):
    _, corpus, principal = env
    wb = Workbook()
    ws = wb.active
    ws.append(["item", "qty"])
    ws.append(["gasket", 4])
    ws.append(["bearing", 6])
    buf = io.BytesIO()
    wb.save(buf)
    workspace.write_bytes("eng-001", "stock.xlsx", buf.getvalue(), origin="upload")

    ctx = ToolContext(principal=principal, corpus=corpus, agent=None)
    code = (
        "import csv\n"
        "rows = list(csv.reader(open('stock.csv')))[1:]\n"
        "total = sum(int(r[1]) for r in rows)\n"
        "print('Step 1: total =', total)\n"
        "open('total.txt', 'w').write(str(total))\n"
    )
    ok, observation, _ = call_tool(ctx, "run_python", {"code": code, "input_files": ["stock.xlsx"]})
    result = json.loads(observation)
    assert ok, observation
    assert "total = 10" in result["stdout"]
    assert result["files_saved_to_workspace"] == ["total.txt"]
    assert workspace.read_bytes("eng-001", "total.txt") == b"10"


def test_write_spreadsheet_keeps_numbers_numeric(env):
    _, corpus, principal = env
    ctx = ToolContext(principal=principal, corpus=corpus, agent=None)
    ctx.observe_numbers("Jan consumption was 1200 t")
    ok, _, _ = call_tool(ctx, "write_spreadsheet", {"filename": "q1.xlsx", "rows": [["Month", "t"], ["Jan", "1200"]]})
    assert ok
    ws = load_workbook(io.BytesIO(workspace.read_bytes("eng-001", "q1.xlsx"))).active
    assert ws["A1"].value.startswith("CLASSIFICATION:")
    assert ws["B3"].value == 1200


def test_repeated_identical_call_is_refused(env):
    db, corpus, principal = env
    call = {"tool": "list_files", "args": {}}
    agent = ScriptedAgent([{"plan": []}, call, call, call, {"final_answer": "ok", "citations": []}])
    res = run_agent_task("list my files", principal, agent, db, corpus)
    assert [s.ok for s in res.agent_trace] == [True, True, False]
    assert "already made this exact call" in res.agent_trace[-1].observation


def test_unknown_tool_is_reported_not_raised(env):
    _, corpus, principal = env
    ctx = ToolContext(principal=principal, corpus=corpus, agent=None)
    ok, observation, _ = call_tool(ctx, "rm_rf", {})
    assert not ok and "Unknown tool" in observation


def test_spreadsheet_rejects_numbers_not_seen_anywhere(env):
    _, corpus, principal = env
    ctx = ToolContext(principal=principal, corpus=corpus, agent=None)
    ctx.observe_numbers("Jan 1200, Feb 1350")
    ok, observation, _ = call_tool(ctx, "write_spreadsheet", {
        "filename": "q1.xlsx", "rows": [["Month", "t"], ["Jan", 1200], ["Feb", 1350], ["Total", 2600]]})
    assert not ok and "2600" in observation  # invented total (really 2550) is refused
    ok, _, _ = call_tool(ctx, "write_spreadsheet", {
        "filename": "q1.xlsx", "rows": [{"Month": "Jan", "t": 1200}, {"Month": "Feb", "t": 1350}]})
    assert ok
    ws = load_workbook(io.BytesIO(workspace.read_bytes("eng-001", "q1.xlsx"))).active
    assert [c.value for c in ws[2]] == ["Month", "t"]  # dict rows -> header row


def test_sandbox_output_must_be_copied_not_retyped(env):
    _, corpus, principal = env
    workspace.write_bytes("eng-001", "stock.csv", b"item,qty,min\ngasket,4,10\nseal,9,5\n", origin="upload")
    ctx = ToolContext(principal=principal, corpus=corpus, agent=None)
    code = "low = [dict(item=r['item'], short=r['min'] - r['qty']) for r in read_rows('stock.csv') if r['qty'] < r['min']]\nwrite_rows('low.csv', low)\nprint(low)"
    ok, observation, _ = call_tool(ctx, "run_python", {"code": code})  # input_files omitted on purpose
    assert ok, observation
    ok, observation, _ = call_tool(ctx, "write_spreadsheet", {"filename": "low.xlsx", "rows": [["item", "short"], ["gasket", 6]]})
    assert not ok and "from_csv='low.csv'" in observation
    ok, _, _ = call_tool(ctx, "write_spreadsheet", {"filename": "low.xlsx", "from_csv": "low.csv"})
    assert ok
    ws = load_workbook(io.BytesIO(workspace.read_bytes("eng-001", "low.xlsx"))).active
    assert [[c.value for c in r] for r in ws.iter_rows(min_row=2)] == [["item", "short"], ["gasket", 6]]


def test_model_server_failures_stop_the_run_honestly(env, monkeypatch):
    db, corpus, principal = env
    monkeypatch.setattr("harness.agent_loop.time.sleep", lambda s: None)

    class Broken(ScriptedAgent):
        def chat_json(self, messages, purpose="planning"):
            raise RuntimeError("model runner stopped")

    res = run_agent_task("anything", principal, Broken([]), db, corpus)
    assert res.status == "abstained"
    assert "failed 3 times" in res.answer


@pytest.mark.parametrize("q,expected", [
    ("draft a Word approval note gas_leak_note.docx", {"docx"}),
    ("Save them to reorder_list.xlsx with the shortfall", {"xlsx"}),
    ("make a 5-slide PowerPoint on the gas leak procedure", {"pptx"}),
    ("what does the word 'isolation' mean in the SOP?", set()),
])
def test_requested_deliverables(q, expected):
    from harness.agent_loop import requested_deliverables
    assert requested_deliverables(q) == expected


def test_agent_cannot_finish_without_requested_file(env):
    db, corpus, principal = env
    agent = ScriptedAgent([
        {"plan": ["1. write the note"]},
        {"final_answer": "Here is your note: isolate and evacuate.", "citations": []},  # skipped the file
        {"tool": "create_word_document", "args": {"filename": "note.docx", "title": "Note", "content": "Isolate."}},
        {"final_answer": "Created note.docx.", "citations": []},
    ])
    res = run_agent_task("Draft a Word note note.docx on isolation", principal, agent, db, corpus)
    assert res.status == "answered"
    assert [a.name for a in res.artifacts] == ["note.docx"]
    assert any("NOT DONE" in c for c in agent.calls)


def test_missing_deliverable_is_never_reported_as_answered(env):
    db, corpus, principal = env
    stubborn = {"final_answer": "", "citations": []}
    agent = ScriptedAgent([{"plan": []}] + [stubborn] * 8)
    res = run_agent_task("Draft a Word note note.docx on isolation", principal, agent, db, corpus)
    assert res.status == "abstained"
    assert "note.docx" not in [a.name for a in res.artifacts]
    assert ".docx file was not produced" in res.answer
